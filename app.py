from flask import Flask, render_template, request, jsonify, send_file

from PIL import Image, ImageOps, ImageFilter
import io
import os
import base64
import tempfile

app = Flask(__name__)

# ============================================================
# CONFIGURATION
# ============================================================

MODEL_PATH = os.path.join("models", "doodle_cnn.keras")

CLASS_NAMES = [
    "apple",
    "cat",
    "flower",
    "house",
    "star",
    "tree"
]

model = None
ocr_reader = None


# ============================================================
# LOAD DOODLE CNN
# ============================================================

def load_model():
    global model

    if not os.path.exists(MODEL_PATH):
        print("⚠️ Doodle CNN model not found.")
        return

    try:
        import tensorflow as tf

        model = tf.keras.models.load_model(MODEL_PATH)

        print("✅ Doodle CNN loaded successfully.")

    except Exception as e:
        model = None
        print("⚠️ Could not load doodle CNN:")
        print(e)


load_model()


# ============================================================
# LOAD OCR
# ============================================================

def load_ocr():
    global ocr_reader

    try:
        import easyocr

        print("🧠 Loading EasyOCR...")

        ocr_reader = easyocr.Reader(
            ["en"],
            gpu=False,
            verbose=False
        )

        print("✅ EasyOCR loaded successfully.")

    except Exception as e:
        ocr_reader = None

        print("⚠️ EasyOCR could not be loaded:")
        print(e)


# Load OCR once when the server starts.
load_ocr()


# ============================================================
# DECODE CANVAS IMAGE
# ============================================================

def decode_canvas(data_url):

    if not data_url:
        raise ValueError("No canvas image received.")

    if "," not in data_url:
        raise ValueError("Invalid canvas image.")

    try:
        raw = base64.b64decode(
            data_url.split(",", 1)[1]
        )

        image = Image.open(
            io.BytesIO(raw)
        )

        return image.convert("RGBA")

    except Exception as e:
        raise ValueError(
            f"Could not decode canvas image: {e}"
        )


# ============================================================
# PREPARE IMAGE FOR DOODLE CNN
# ============================================================

def prepare_doodle_image(image):

    # Create a black background.
    background = Image.new(
        "L",
        image.size,
        0
    )

    # Extract alpha channel.
    alpha = image.getchannel("A")

    # Convert visible strokes into grayscale.
    gray = image.convert("L")

    # Keep only visible pixels.
    background.paste(
        gray,
        mask=alpha
    )

    # Resize for CNN.
    background = background.resize(
        (28, 28),
        Image.Resampling.LANCZOS
    )

    return background


# ============================================================
# HOME
# ============================================================

@app.route("/")
def index():
    return render_template("index.html")


# ============================================================
# DOODLE CNN
# ============================================================

@app.post("/predict-doodle")
def predict_doodle():

    try:

        data = request.get_json(silent=True) or {}

        img = decode_canvas(
            data.get("image")
        )

        if model is None:

            return jsonify({
                "ok": False,
                "message":
                    "CNN model is not available. "
                    "Check models/doodle_cnn.keras."
            }), 503

        import numpy as np

        img = prepare_doodle_image(img)

        arr = np.array(
            img,
            dtype="float32"
        ) / 255.0

        arr = arr.reshape(
            1,
            28,
            28,
            1
        )

        probs = model.predict(
            arr,
            verbose=0
        )[0]

        idx = int(
            np.argmax(probs)
        )

        scores = {
            CLASS_NAMES[i]:
                round(float(probs[i]) * 100, 1)

            for i in range(
                len(CLASS_NAMES)
            )
        }

        return jsonify({

            "ok": True,

            "prediction":
                CLASS_NAMES[idx],

            "confidence":
                round(
                    float(probs[idx]) * 100,
                    1
                ),

            "scores": scores

        })

    except Exception as e:

        print("❌ CNN ERROR:", e)

        return jsonify({
            "ok": False,
            "message": str(e)
        }), 400


# ============================================================
# SAVE PNG
# ============================================================

@app.post("/save-doodle")
def save_doodle():

    try:

        data = request.get_json(silent=True) or {}

        img = decode_canvas(
            data.get("image")
        )

        # White background.
        background = Image.new(
            "RGBA",
            img.size,
            (248, 255, 252, 255)
        )

        # Use alpha as mask so only the drawing is placed.
        background.alpha_composite(img)

        fd, path = tempfile.mkstemp(
            suffix=".png"
        )

        os.close(fd)

        background.save(
            path,
            "PNG"
        )

        return send_file(
            path,
            as_attachment=True,
            download_name="neurocanvas-doodle.png"
        )

    except Exception as e:

        print("❌ PNG ERROR:", e)

        return jsonify({
            "ok": False,
            "message": str(e)
        }), 400


# ============================================================
# PREPARE HANDWRITING FOR OCR
# ============================================================

def prepare_ocr_image(image):

    # Create a clean white background.
    white = Image.new(
        "L",
        image.size,
        255
    )

    # Extract alpha channel.
    alpha = image.getchannel("A")

    # Convert strokes to grayscale.
    gray = image.convert("L")

    # Because the canvas uses light strokes,
    # convert visible strokes into dark strokes.
    inverted = ImageOps.invert(gray)

    # Put visible strokes onto white background.
    white.paste(
        inverted,
        mask=alpha
    )

    # Improve contrast.
    white = ImageOps.autocontrast(
        white
    )

    # Slight smoothing.
    white = white.filter(
        ImageFilter.MedianFilter(3)
    )

    # Add a little border.
    padding = 30

    canvas = Image.new(
        "L",
        (
            white.width + padding * 2,
            white.height + padding * 2
        ),
        255
    )

    canvas.paste(
        white,
        (padding, padding)
    )

    return canvas


# ============================================================
# HANDWRITING OCR
# ============================================================

@app.post("/ocr")
def ocr():

    try:

        print("📝 OCR request received.")

        data = request.get_json(
            silent=True
        ) or {}

        if not data.get("image"):

            return jsonify({
                "ok": False,
                "message":
                    "No handwriting image received."
            }), 400

        # Make sure OCR exists.
        if ocr_reader is None:

            return jsonify({
                "ok": False,
                "message":
                    "EasyOCR is not available. "
                    "Run: pip install easyocr"
            }), 503

        # Decode canvas.
        img = decode_canvas(
            data.get("image")
        )

        # Prepare for OCR.
        processed = prepare_ocr_image(
            img
        )

        import numpy as np

        image_array = np.array(
            processed
        )

        print("🧠 Running neural OCR...")

        results = ocr_reader.readtext(
            image_array,
            detail=0,
            paragraph=True,
            text_threshold=0.4,
            low_text=0.2,
            link_threshold=0.2
        )

        # Convert OCR results into text.
        text = " ".join(
            str(x).strip()
            for x in results
            if str(x).strip()
        )

        if not text:

            text = "No handwriting detected."

        print("✅ OCR RESULT:", text)

        return jsonify({
            "ok": True,
            "text": text
        })

    except Exception as e:

        print("❌ OCR ERROR:")
        print(e)

        return jsonify({
            "ok": False,
            "message": str(e)
        }), 500


# ============================================================
# EXPORT DOCX
# ============================================================

@app.post("/export")
def export_doc():

    try:

        data = request.get_json(
            silent=True
        ) or {}

        text = data.get(
            "text",
            ""
        ).strip()

        if not text:

            return jsonify({
                "ok": False,
                "message":
                    "There is no text to export."
            }), 400

        from docx import Document

        doc = Document()

        doc.add_heading(
            "Canvaru Document",
            0
        )

        doc.add_paragraph(
            text
        )

        fd, path = tempfile.mkstemp(
            suffix=".docx"
        )

        os.close(fd)

        doc.save(path)

        return send_file(
            path,
            as_attachment=True,
            download_name=
                "neurocanvas-document.docx"
        )

    except Exception as e:

        print("❌ EXPORT ERROR:", e)

        return jsonify({
            "ok": False,
            "message": str(e)
        }), 400


# ============================================================
# RUN SERVER
# ============================================================

if __name__ == "__main__":

    print("")
    print("====================================")
    print("     CANVARU NEURAL CREATIVE STUDIO")
    print("====================================")
    print("")

    if model:
        print("✅ CNN: ONLINE")
    else:
        print("⚠️ CNN: NOT AVAILABLE")

    if ocr_reader:
        print("✅ OCR: ONLINE")
    else:
        print("⚠️ OCR: NOT AVAILABLE")

    print("")
    print("🌐 http://127.0.0.1:5000")
    print("")

    app.run(
        debug=True,
        port=5000
    )
    
    
    